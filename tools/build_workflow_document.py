from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(__file__).resolve().parents[1] / "AI_Due_Diligence_Copilot_Workflow.docx"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = "D9D9D9") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), color)


def style_cell(cell, bold: bool = False, color: str | None = None, size: int = 9) -> None:
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_border(cell)
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.space_before = Pt(2)
        for run in paragraph.runs:
            run.font.name = "Aptos"
            run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
            run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
            run.font.size = Pt(size)
            run.bold = bold
            if color:
                run.font.color.rgb = RGBColor.from_string(color)


def table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.autofit = False
    header = t.rows[0]
    for i, text in enumerate(headers):
        header.cells[i].width = Inches(widths[i])
        header.cells[i].text = text
        set_cell_shading(header.cells[i], "1F4E79")
        style_cell(header.cells[i], bold=True, color="FFFFFF")
    for row_index, data in enumerate(rows):
        cells = t.add_row().cells
        for i, text in enumerate(data):
            cells[i].width = Inches(widths[i])
            cells[i].text = text
            if row_index % 2:
                set_cell_shading(cells[i], "F2F6FA")
            style_cell(cells[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(3)


def bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.add_run(text)


def heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(text, level=level)
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(5)
    for run in p.runs:
        run.font.color.rgb = RGBColor(0, 0, 0)


def body(doc: Document, text: str) -> None:
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.12


def main() -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    styles["Normal"].font.size = Pt(10.5)
    styles["Title"].font.name = "Aptos Display"
    styles["Title"].font.size = Pt(24)
    styles["Title"].font.color.rgb = RGBColor(0, 0, 0)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.add_run("AI Due Diligence Copilot Workflow")
    subtitle = doc.add_paragraph("System workflow frontend backend connection and operating guide")
    subtitle.paragraph_format.space_after = Pt(16)
    for run in subtitle.runs:
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(80, 80, 80)

    body(doc, "This guide explains how the local-first AI Due Diligence Copilot moves data from a browser upload to an evidence-backed research response, then stores that work for later review. It also records the verification performed on October 31 2026. The working stack has an authenticated Next.js frontend, a FastAPI service, user-isolated SQLite and Chroma workspaces, and local Ollama models for embeddings and generation.")
    body(doc, "Conclusion: the application builds successfully, its backend test suite passes, its authenticated workspace endpoints return data, and the sign-in UI renders locally. The only environment action required before ordinary use is keeping the Ollama service running with the configured models. The development environment used for this review prevents Python subprocesses from reaching the host Ollama socket, so the final LLM generation request could not be executed from that sandbox even though the Ollama HTTP endpoint itself is reachable and the models are installed.")

    heading(doc, "System at a Glance")
    table(doc, ["Layer", "Responsibility", "Key implementation"], [
        ["Browser frontend", "Authentication, company and document workflow, research chat, charts, reports and settings", "Next.js 16, React 19, TypeScript, Tailwind, Recharts"],
        ["API", "JWT session validation, user-scoped endpoints, uploads, background indexing, research orchestration", "FastAPI in dd_copilot api.py"],
        ["Workspace data", "Companies, document metadata, chunks, facts, research history and reports", "One SQLite database plus Chroma directory per user"],
        ["Ingestion", "Streaming file save, parsing, chunking, fact extraction, embeddings and vector persistence", "ingest.py, loaders.py, dd_copilot ingestion"],
        ["AI and retrieval", "Hybrid vector and BM25 evidence retrieval, calculations, answer generation and citation validation", "CopilotPipeline, Chroma, rank bm25, Ollama"],
    ], [1.2, 2.75, 2.45])

    heading(doc, "End to End Workflow")
    body(doc, "The normal user journey is a continuous path rather than separate products. A user logs in, creates or selects a company, uploads a source document, waits for indexing, asks a question, and opens the saved answer or report. Each stage is scoped to the authenticated user so one workspace cannot query another workspace's files or vectors.")
    table(doc, ["Step", "What happens", "Visible frontend result", "Backend work"], [
        ["1 Login", "The user submits email and password.", "The app shell loads after a successful sign-in.", "POST api auth login validates PBKDF2 password hash and returns a signed bearer token."],
        ["2 Workspace", "The browser requests companies, documents, research and dashboard data using the token.", "Dashboard, company pages, watchlist, documents and reports populate from the current workspace.", "FastAPI resolves the user, initializes that user's workspace, then queries the user SQLite database."],
        ["3 Upload", "The user selects a company, type and fiscal year, then chooses a supported file.", "Documents page shows a queued job and polls until it is indexed or failed.", "The upload is streamed in 8 MB chunks to local workspace storage and checked against the configured file limit."],
        ["4 Index", "The background job parses the file and derives chunks and financial facts.", "Document row changes to indexed with chunk count; company views gain financial and document data.", "Text is chunked, persisted in SQLite, embedded with mxbai embed large and stored in Chroma."],
        ["5 Research chat", "The user asks a question globally or for a selected company.", "A chat turn displays an answer, confidence, citations and evidence cards.", "Hybrid retrieval combines vector and BM25 candidates, ranks evidence, computes fact-based metrics and invokes qwen3."],
        ["6 Persist", "The completed question can be revisited later.", "Research chat history and Reports show the saved analysis.", "Research result, evidence pack and a report with an analysis section are stored in SQLite."],
    ], [0.65, 1.8, 2.0, 1.95])

    heading(doc, "Frontend to Backend Contract")
    body(doc, "The client adapter in frontend lib api client ts is the single browser-facing API layer. It obtains the base URL from NEXT PUBLIC API URL, adds the bearer token from local storage, converts server payloads into page types, and sends multipart form data only for file uploads. All other app pages use this adapter rather than embedding API calls independently.")
    table(doc, ["Frontend action", "HTTP endpoint", "Server response used by the UI"], [
        ["Sign up and sign in", "POST api auth register and POST api auth login", "User profile and token"],
        ["Session shell and settings", "GET api auth me and POST api auth logout", "Authenticated user record or cleared browser token"],
        ["Companies and dashboard", "GET and POST api companies, GET api companies id, GET api dashboard", "Company data, documents, facts, segments, geography and recent research"],
        ["Documents", "GET api documents, POST api documents upload, GET api documents jobs id", "Document rows plus job status and progress"],
        ["Research chat and history", "POST ask, GET api research, GET api research id", "Answer, evidence pack, citations, confidence, timings and saved research"],
        ["Reports", "GET api reports and GET api reports id", "Saved report summaries and report sections"],
    ], [1.55, 2.35, 2.5])

    heading(doc, "Document Handling and Large File Support")
    body(doc, "The upload route accepts CSV, PDF, TXT, DOCX, Markdown, JSON, XLSX and XLS files. It saves each incoming request in streamed 8 MB pieces rather than reading the whole document into memory. The default maximum is 2 GB, controlled by RAG MAX FILE SIZE MB. The system also applies PDF page, Excel row, JSON record, directory depth and file-count limits to prevent one source file from exhausting local resources.")
    bullet(doc, "Files are stored under workspaces user id data company folder. Original data is not mixed between users.")
    bullet(doc, "Each upload has a DocumentJob record with queued, processing, indexed or failed status. The Documents UI can poll this record and show progress.")
    bullet(doc, "Indexing refreshes the user's cached research pipeline so a new upload becomes eligible for the next chat question.")
    bullet(doc, "Text chunks, financial facts and metadata are linked back to the originating document so the UI can show filenames, pages, sections and evidence excerpts.")

    heading(doc, "Retrieval and Answer Generation")
    body(doc, "For a question, the API classifies the query, retrieves relevant chunks with vector and BM25 search, merges and reranks the candidates, loads eligible financial facts, and calculates derived metrics. The prompt is assembled from that evidence pack and sent to the local qwen3 model. The response is checked against the citations before it is returned. When no evidence is available, the pipeline returns a structured insufficient-evidence answer instead of inventing a conclusion.")
    table(doc, ["Control", "Configured value", "Purpose"], [
        ["Embedding model", "mxbai embed large", "Creates local vectors for document chunks"],
        ["Chat model", "qwen3", "Generates the research analysis from retrieved evidence"],
        ["Vector and BM25 candidates", "20 and 20", "Provides diverse retrieval candidates before fusion"],
        ["Final reranked chunks", "8", "Bounds prompt size while keeping the strongest evidence"],
        ["LLM timeout", "180 seconds", "Allows local model responses while avoiding an indefinitely waiting API call"],
        ["Upload limit", "2048 MB", "Supports large source documents through streaming upload"],
    ], [1.85, 1.35, 3.2])

    heading(doc, "Security and Data Boundaries")
    body(doc, "Identity data is kept in the authentication database. Each authenticated user has a separate local workspace containing their SQLite database, Chroma collection, uploads, chunks, extracted facts, research and reports. Authorization is enforced by the bearer token dependency on every workspace endpoint. Passwords use PBKDF2 SHA 256 with a per-password salt and 310000 iterations. Tokens are signed with DD AUTH SECRET and expire according to DD AUTH TOKEN TTL.")
    body(doc, "For a production deployment, replace the local development secret and default administrator password, keep the authentication database central if needed, use HTTPS, restrict CORS to the deployed frontend origin, and expose the local workspace service only through the intended network boundary.")

    heading(doc, "Startup and Operating Procedure")
    table(doc, ["Order", "Command or action", "Expected result"], [
        ["1", "Start Ollama and confirm qwen3 and mxbai embed large are installed", "Local AI service answers on 127.0.0.1 port 11434"],
        ["2", "Activate .venv and install requirements txt", "FastAPI and upload dependency python multipart are available"],
        ["3", "Run uvicorn dd copilot api app on port 8000", "Health endpoint returns status ok"],
        ["4", "Run npm run dev in frontend", "Next.js serves the application on port 3000"],
        ["5", "Sign in, add a company if needed and upload a supported source", "Upload job progresses to indexed"],
        ["6", "Open Research Chat, select the company and ask an evidence question", "Answer, citations and persisted history appear"],
    ], [0.55, 3.35, 2.5])

    heading(doc, "Configuration and Required Services")
    body(doc, "No OpenAI, Anthropic or other hosted AI API keys are required. This project is configured for local Ollama. The necessary model service and model files are present: qwen3, mxbai embed large and llama3 2. The review corrected OLLAMA BASE URL to http 127 0 0 1 11434 so the application targets the local IPv4 service explicitly.")
    table(doc, ["Setting", "Required for", "Status during review"], [
        ["DD AUTH SECRET", "Signing local sessions", "Set; replace with a long random secret before production"],
        ["DD ADMIN EMAIL and DD ADMIN PASSWORD", "Initial local administrator seed", "Set and authenticated successfully"],
        ["OLLAMA BASE URL", "Embeddings and chat generation", "Set to the local IPv4 service"],
        ["qwen3 and mxbai embed large", "Generation and embeddings", "Installed in local Ollama"],
        ["NEXT PUBLIC API URL", "Browser to FastAPI connection", "Configured for local port 8000"],
        ["python multipart", "Multipart document uploads", "Installed from requirements during review"],
    ], [1.85, 2.55, 2.0])

    heading(doc, "Verification Record")
    table(doc, ["Check", "Result", "Notes"], [
        ["Backend tests", "Passed", "48 tests passed across the copilot and financial reasoning suites"],
        ["FastAPI import", "Passed", "21 application routes loaded after installing python multipart"],
        ["Authenticated APIs", "Passed", "Login, profile, companies, documents, dashboard, research and reports returned expected authenticated responses"],
        ["Existing workspace", "Passed", "1 company, 1 document, 228 chunks and 84 financial facts were available to the admin workspace"],
        ["Frontend lint", "Passed with warnings", "No errors; warnings are limited to unused template imports and recommended React navigation and effect refinements"],
        ["Frontend production build", "Passed", "Next.js Webpack production build completed with TypeScript and all 18 routes generated"],
        ["Frontend visual check", "Passed", "Local sign-in page rendered in the browser with expected authentication controls"],
        ["Live Ollama generation", "Sandbox constrained", "Ollama HTTP tags endpoint and required models were reachable; Python subprocess network isolation in the review environment blocked the final model call"],
    ], [1.6, 1.25, 3.55])

    heading(doc, "Recommended Next Actions")
    bullet(doc, "Change the default administrator password and DD AUTH SECRET before any shared or production deployment.")
    bullet(doc, "Run one Research Chat question from the normal desktop terminal after startup to confirm the local Python process can contact Ollama outside sandboxed review tooling.")
    bullet(doc, "Keep the frontend lint warnings as a short cleanup task. They do not block the compiled application, but replacing internal window navigation with the Next router and tightening a few effects will improve maintainability.")
    bullet(doc, "If deploying beyond localhost, set a production frontend origin in CORS, use HTTPS and choose a persistent database and volume strategy for the authentication store and every user workspace.")

    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run("AI Due Diligence Copilot   |   Workflow and operating guide")
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(100, 100, 100)

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
